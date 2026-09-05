"""
Word (.docx) Export Module
Generates clean, structured Word documents for meeting minutes using python-docx.
"""

import io
import os
import re
from typing import Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from src.utils.sanitizer import sanitize_meeting_data

class WordExporter:
    def __init__(self, output_dir: str = "outputs/word"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate(self, meeting_data: Dict[str, Any], filename: Optional[str] = None) -> bytes:
        # Strictly sanitize meeting data so no secrets or API keys can enter DOCX
        meeting_data = sanitize_meeting_data(meeting_data)
        doc = docx.Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(meeting_data.get("title", "Official Meeting Minutes"))
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(30, 58, 138) # #1e3a8a
        p_title.paragraph_format.space_after = Pt(4)

        # Subtitle / Metadata
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(12)
        run_meta = p_meta.add_run(
            f"Date: {meeting_data.get('date', 'N/A')}   |   "
            f"Time: {meeting_data.get('time', 'Standard Session') or 'Standard Session'}   |   "
            f"Status: {meeting_data.get('status', 'Draft')}"
        )
        run_meta.font.size = Pt(10)
        run_meta.font.color.rgb = RGBColor(100, 116, 139)

        # Attendees
        attendees = meeting_data.get('attendees', [])
        attendees_str = ", ".join(attendees) if isinstance(attendees, list) else str(attendees)
        p_att = doc.add_paragraph()
        p_att.paragraph_format.space_after = Pt(14)
        run_att_label = p_att.add_run("Attendees: ")
        run_att_label.bold = True
        run_att_label.font.size = Pt(10.5)
        run_att_val = p_att.add_run(attendees_str or "All Invited Members")
        run_att_val.font.size = Pt(10.5)

        # 1. Executive Summary
        self._add_heading(doc, "1. Executive Summary")
        summary_text = meeting_data.get("executive_summary", "No summary available.")
        p_sum = doc.add_paragraph(summary_text)
        p_sum.paragraph_format.space_after = Pt(12)

        # 2. Key Discussion Points
        discussions = meeting_data.get("discussion_points", [])
        if discussions:
            self._add_heading(doc, "2. Key Discussion Points")
            for pt in discussions:
                doc.add_paragraph(pt, style='List Bullet')
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 3. Decisions Made
        decisions = meeting_data.get("decisions", [])
        if decisions:
            self._add_heading(doc, "3. Decisions Made & Resolutions")
            for d in decisions:
                p_dec = doc.add_paragraph(style='List Bullet')
                r_check = p_dec.add_run("✔ ")
                r_check.bold = True
                r_check.font.color.rgb = RGBColor(5, 150, 105) # Green
                p_dec.add_run(d)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 4. Action Items Table
        action_items = meeting_data.get("action_items", [])
        self._add_heading(doc, "4. Action Items & Deliverables")

        if action_items:
            table = doc.add_table(rows=1, cols=5)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            # Set column widths
            col_widths = [Inches(2.7), Inches(1.1), Inches(1.1), Inches(0.9), Inches(0.9)]

            hdr_cells = table.rows[0].cells
            headers = ["Task Description", "Owner", "Deadline", "Priority", "Status"]
            for i, h in enumerate(headers):
                hdr_cells[i].text = h
                hdr_cells[i].width = col_widths[i]
                # Header styling
                shading_elm = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
                hdr_cells[i]._tc.get_or_add_tcPr().append(shading_elm)
                for run in hdr_cells[i].paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)

            for item in action_items:
                row_cells = table.add_row().cells
                vals = [
                    item.get('task', ''),
                    item.get('owner', 'Unassigned'),
                    item.get('deadline', 'TBD'),
                    item.get('priority', 'Medium'),
                    item.get('status', 'Pending')
                ]
                for i, v in enumerate(vals):
                    row_cells[i].text = str(v)
                    row_cells[i].width = col_widths[i]
                    for run in row_cells[i].paragraphs[0].runs:
                        run.font.size = Pt(9)
        else:
            doc.add_paragraph("No specific action items recorded.")

        # Footer
        doc.add_paragraph().paragraph_format.space_after = Pt(16)
        p_foot = doc.add_paragraph()
        run_foot = p_foot.add_run(
            "Generated by AI-Based Meeting Minutes Generator | Academic Project S.Y. B.Sc. AI & ML (2026-2027) | Guide: Prof. Sarita Byagar"
        )
        run_foot.font.size = Pt(8)
        run_foot.font.italic = True
        run_foot.font.color.rgb = RGBColor(148, 163, 184)

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()

        # Save to outputs/word/
        safe_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', meeting_data.get('title', 'meeting_minutes'))[:40]
        meeting_id = meeting_data.get('id', 'temp')
        file_path = os.path.join(self.output_dir, filename or f"{safe_title}_{meeting_id}.docx")
        with open(file_path, "wb") as f:
            f.write(docx_bytes)

        return docx_bytes

    def _add_heading(self, doc: docx.Document, text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
